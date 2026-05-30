import openpyxl

from app.files.service import get_file_if_visible, open_file
from app.tools.write_xlsx import write_xlsx


def test_write_xlsx_creates_multi_sheet(client, member_token, env):
    from app.config import get_settings
    from app.deps import _session_factory
    from app.models import User
    db = _session_factory()
    settings = get_settings()
    user = db.query(User).filter_by(username="alice").one()
    sheets = {
        "S1": [["a", "b"], [1, 2]],
        "S2": [["only"]],
    }
    out = write_xlsx(name="out.xlsx", sheets=sheets,
                     db=db, settings=settings, user_id=user.id, user_role="member")
    row = get_file_if_visible(db, out["file_id"], user.id, "member")
    assert row is not None

    wb = openpyxl.load_workbook(str(open_file(settings, row)), read_only=True)
    assert set(wb.sheetnames) == {"S1", "S2"}
    s1 = [list(r) for r in wb["S1"].iter_rows(values_only=True)]
    assert s1 == [["a", "b"], [1, 2]]
