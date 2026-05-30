import io

import docx

from app.tools._helpers import write_personal
from app.tools.registry import register_tool


def _markdown_to_docx(md: str) -> bytes:
    """Minimal markdown rendering: headings, paragraphs, bullets, numbered."""
    d = docx.Document()
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line:
            d.add_paragraph("")
            continue
        if line.startswith("### "):
            d.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            d.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            d.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            d.add_paragraph(line[2:], style="List Bullet")
        elif len(line) > 2 and line[0].isdigit() and line[1:3] in (". ", ") "):
            d.add_paragraph(line[3:], style="List Number")
        else:
            d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@register_tool(
    name="write_docx",
    description=(
        "Create a .docx Word file from markdown content (headings #/##/###, "
        "bullets, numbered lists, paragraphs). Saved to the user's personal workspace."
    ),
    parameters={
        "type": "object",
        "properties": {
            "name": {"type": "string"},
            "markdown": {"type": "string", "description": "Markdown source"},
        },
        "required": ["name", "markdown"],
    },
)
def write_docx(name: str, markdown: str, *, db, settings, user_id: int, user_role: str) -> dict:
    content = _markdown_to_docx(markdown)
    row = write_personal(
        db, settings, user_id=user_id, name=name,
        content_bytes=content,
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    return {"file_id": row.id, "name": row.original_name}
